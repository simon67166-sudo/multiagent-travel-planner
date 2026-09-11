"""
把组员手动整理的 `数据库原始.xlsx`（港澳达人帖子，本地人认证 + 小红书两个来源）
转成标准化 JSON 数据库：`orchestrator/data_sources/hk_macau_posts.json`。

只负责"读原始表格 -> 标准化字段 + 编号"，不碰 Chroma —— 真正灌进向量库是
`import_hk_macau_data.py` 的事，两步分开方便调试（标准化出错不用重新导入一次库）。

编号规则（城市前缀 + 类型字母 + 三位数字，类型内连续编号，不区分本地人/小红书来源）：
  M=Macau, H=Hong Kong
  F=饮食, S=景点, T=Tips（推荐，按分类正常拆开）
  N=避坑（不细分，统一一个序列）
  例：MF001, MS001, MT001, MN001 / HF001, HS001, HT001, HN001

`verified_local` 字段区分数据来源：来自"本地人XX"表 = True（优先加权用，见 store.py
query_similar_posts 的 boost 逻辑）；来自"小红书XX"表 = False。这个字段是原始数据
自带的真实来源标记，不是靠算法"判断"出来的——组员整理的时候已经分好表了。

原始表格本身字段不统一（不同分类小节的列不一样，比如"本地人香港"的 Tips 小节没有
单独的"名称"列，内容直接写在名称位置），解析时按每个小节自己的表头动态取列，
不能假设固定列位置。
"""

import json
from pathlib import Path

import openpyxl

_ROOT = Path(__file__).resolve().parent.parent
_SRC = _ROOT / "数据库原始.xlsx"
_OUT_DIR = Path(__file__).resolve().parent / "data_sources"
_OUT_FILE = _OUT_DIR / "hk_macau_posts.json"

_CITY_PREFIX = {"澳门": "M", "香港": "H"}
_CATEGORY_LETTER = {"饮食": "F", "景点": "S", "Tips": "T"}


def _find_header_rows(rows: list[list]) -> list[tuple[int, dict[int, str]]]:
    """找到所有 header 行（第0列是 'id' 的行），返回 (行号, {列号: 表头名}) 列表。"""
    headers = []
    for i, r in enumerate(rows):
        if r and r[0] == "id":
            headers.append((i, {idx: v for idx, v in enumerate(r) if v is not None}))
    return headers


def _parse_local_sheet(rows: list[list]) -> tuple[list[dict], list[dict]]:
    """本地人澳门 / 本地人香港：左表推荐（按分类分小节），右表避坑（只在第一个 header 出现过）。"""
    headers = _find_header_rows(rows)
    recommend, avoid = [], []

    for h_idx, (row_i, cols) in enumerate(headers):
        next_row = headers[h_idx + 1][0] if h_idx + 1 < len(headers) else len(rows)
        right_id_col = next((idx for idx in sorted(cols) if idx > 0 and cols[idx] == "id"), None)
        left_cols = {name: idx for idx, name in cols.items() if right_id_col is None or idx < right_id_col}
        right_cols = {name: idx for idx, name in cols.items() if right_id_col is not None and idx >= right_id_col}

        for r in rows[row_i + 1 : next_row]:
            if not r or all(v is None for v in r):
                continue

            def g(name, cols=left_cols, row=r):
                idx = cols.get(name)
                return row[idx] if idx is not None and idx < len(row) else None

            if g("分类") in ("饮食", "景点", "Tips"):
                recommend.append(
                    {
                        "category": g("分类"),
                        "name": g("名称"),
                        "address": g("地址"),
                        "recommender": g("推荐人"),
                        "reason": g("推荐理由"),
                        "tags": [t for t in (g("标签一"), g("标签二"), g("标签三")) if t],
                        "avg_cost": g("人均"),
                        "rating": g("评分"),
                    }
                )

            if right_cols:

                def gr(name, cols=right_cols, row=r):
                    idx = cols.get(name)
                    return row[idx] if idx is not None and idx < len(row) else None

                if gr("名称") or gr("避雷理由"):
                    avoid.append({"name": gr("名称"), "address": gr("地址"), "reason": gr("避雷理由")})

    return recommend, avoid


def _parse_xhs_sheet(rows: list[list]) -> list[dict]:
    """小红书澳门 / 小红书香港：单表，褒贬都有，按分类分小节。"""
    headers = _find_header_rows(rows)
    posts = []
    for h_idx, (row_i, cols) in enumerate(headers):
        next_row = headers[h_idx + 1][0] if h_idx + 1 < len(headers) else len(rows)
        col_by_name = {name: idx for idx, name in cols.items()}
        for r in rows[row_i + 1 : next_row]:
            if not r or all(v is None for v in r):
                continue

            def g(name):
                idx = col_by_name.get(name)
                return r[idx] if idx is not None and idx < len(r) else None

            if g("分类") in ("饮食", "景点", "Tips") and g("名称"):
                posts.append(
                    {
                        "category": g("分类"),
                        "name": g("名称"),
                        "content": g("相关内容"),
                        "likes": g("点赞数"),
                        "recency": g("新/旧(半年以内算新)"),
                        "avg_cost": g("人均"),
                        "rating": g("评分"),
                    }
                )
    return posts


def _next_id(counters: dict[str, int], key: str) -> str:
    counters[key] = counters.get(key, 0) + 1
    return f"{key}{counters[key]:03d}"


def standardize() -> list[dict]:
    wb = openpyxl.load_workbook(_SRC, data_only=True, read_only=True)
    counters: dict[str, int] = {}
    records: list[dict] = []

    for city, local_sheet, xhs_sheet in (("澳门", "本地人澳门", "小红书澳门"), ("香港", "本地人香港", "小红书香港")):
        prefix = _CITY_PREFIX[city]

        local_rows = [list(r) for r in wb[local_sheet].iter_rows(values_only=True)]
        recommend, avoid = _parse_local_sheet(local_rows)
        for item in recommend:
            letter = _CATEGORY_LETTER[item["category"]]
            records.append(
                {
                    "id": _next_id(counters, f"{prefix}{letter}"),
                    "city": city,
                    "category": item["category"],
                    "post_type": "recommend",
                    "verified_local": True,
                    "place": item["name"],
                    "address": item["address"],
                    "caption": item["reason"],
                    "recommender": item["recommender"],
                    "tags": item["tags"],
                    "avg_cost": item["avg_cost"],
                    "rating": item["rating"],
                }
            )
        for item in avoid:
            records.append(
                {
                    "id": _next_id(counters, f"{prefix}N"),
                    "city": city,
                    "category": None,
                    "post_type": "avoid",
                    "verified_local": True,
                    "place": item["name"],
                    "address": item["address"],
                    "avoid_tips": item["reason"],
                }
            )

        xhs_rows = [list(r) for r in wb[xhs_sheet].iter_rows(values_only=True)]
        for item in _parse_xhs_sheet(xhs_rows):
            letter = _CATEGORY_LETTER[item["category"]]
            records.append(
                {
                    "id": _next_id(counters, f"{prefix}{letter}"),
                    "city": city,
                    "category": item["category"],
                    "post_type": "recommend",
                    "verified_local": False,
                    "place": item["name"],
                    "caption": item["content"],
                    "likes": item["likes"],
                    "recency": item["recency"],
                    "avg_cost": item["avg_cost"],
                    "rating": item["rating"],
                }
            )

    return records


def main() -> None:
    records = standardize()
    _OUT_DIR.mkdir(parents=True, exist_ok=True)
    with open(_OUT_FILE, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

    by_prefix: dict[str, int] = {}
    for r in records:
        by_prefix[r["id"][:2]] = by_prefix.get(r["id"][:2], 0) + 1
    print(f"共 {len(records)} 条，写入 {_OUT_FILE}")
    for k in sorted(by_prefix):
        print(f"  {k}: {by_prefix[k]} 条")


if __name__ == "__main__":
    main()
